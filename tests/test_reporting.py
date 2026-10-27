"""Smoke tests for the report generation pipeline (no GUI or network required)."""
import os
import tempfile
import unittest
import zipfile
from types import SimpleNamespace

from openpyxl import load_workbook
import reporting


class ReportingSmokeTests(unittest.TestCase):
    def setUp(self):
        self.app = SimpleNamespace(
            results=[{"ip": "192.168.1.2", "hostname": "host-2", "latency_ms": 4.3,
                      "status": "Excellent", "timestamp": "2026-10-27T10:00:00"}],
            port_results=[{"host": "192.168.1.2", "ip": "192.168.1.2", "port": 443,
                           "service": "https", "state": "open"}],
            dns_results=[{"type": "A", "name": "example.test", "value": "192.0.2.1"}],
            interface_results=[{"name": "Ethernet", "address": "192.168.1.2",
                                "netmask": "255.255.255.0", "status": "up"}],
            monitor_history=[
                {"timestamp": "2026-10-27T10:00:00", "target": "192.168.1.1", "online": True,
                 "latency_ms": 1.4, "line": "sample online"},
                {"timestamp": "2026-10-27T10:00:02", "target": "192.168.1.1", "online": False,
                 "latency_ms": None, "line": "sample lost"},
            ],
            alert_events=[{"timestamp": "2026-10-27T10:00:02", "target": "192.168.1.1",
                           "event": "UNAVAILABLE", "previous_state": "ONLINE",
                           "current_state": "NO RESPONSE", "latency_ms": None}],
            scan_snapshots=[], local_ip_value="192.168.1.2", public_ip_value="Unavailable",
            ip_var=SimpleNamespace(get=lambda: "192.168.1.0/24"),
        )
        self.payload = reporting.build_payload(self.app)

    def test_sample_availability_and_alert_counts(self):
        self.assertEqual(self.payload["summary"]["availability_pct"], 50.0)
        self.assertEqual(self.payload["summary"]["alert_events"], 1)
        self.assertEqual(self.payload["summary"]["packet_loss_pct"], 50.0)

    def test_all_exporters_and_bundle(self):
        exporters = [
            ("technical.pdf", lambda p: reporting.export_pdf(self.payload, p, False)),
            ("executive.pdf", lambda p: reporting.export_pdf(self.payload, p, True)),
            ("workbook.xlsx", lambda p: reporting.export_xlsx(self.payload, p)),
            ("report.html", lambda p: reporting.export_html(self.payload, p)),
            ("report.json", lambda p: reporting.export_json(self.payload, p)),
            ("hosts.csv", lambda p: reporting.export_csv(self.payload, p)),
            ("report.md", lambda p: reporting.export_markdown(self.payload, p)),
            ("report.txt", lambda p: reporting.export_text(self.payload, p)),
            ("bundle.zip", lambda p: reporting.export_bundle(self.payload, p)),
        ]
        with tempfile.TemporaryDirectory() as folder:
            for filename, export in exporters:
                path = os.path.join(folder, filename)
                export(path)
                self.assertGreater(os.path.getsize(path), 0, filename)
            workbook = load_workbook(os.path.join(folder, "workbook.xlsx"), read_only=True)
            self.assertIn("Alert Events", workbook.sheetnames)
            with zipfile.ZipFile(os.path.join(folder, "bundle.zip")) as archive:
                self.assertEqual(len(archive.namelist()), 9)


if __name__ == "__main__":
    unittest.main()
