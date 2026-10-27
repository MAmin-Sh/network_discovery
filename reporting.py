"""Professional multi-format reporting for Network Operations Console.
All report content is generated locally from data collected in the current app session.
"""
from __future__ import annotations
import csv, html, json, os, statistics, textwrap, zipfile
from datetime import datetime
from pathlib import Path

APP_TITLE = "NETWORK OPERATIONS CONSOLE"
NAVY = "14233B"
BLUE = "2457A7"
TEAL = "0F766E"
LIGHT = "EEF3F9"
MUTED = "607086"


def _safe(value):
    return "N/A" if value is None or value == "" else str(value)


def build_payload(app):
    hosts = [dict(r) for r in getattr(app, "results", [])]
    ports = [dict(r) for r in getattr(app, "port_results", [])]
    dns = [dict(r) for r in getattr(app, "dns_results", [])]
    interfaces = [dict(r) for r in getattr(app, "interface_results", [])]
    monitor = [dict(r) for r in getattr(app, "monitor_history", [])]
    now = datetime.now().astimezone()
    # Scanner status labels describe latency quality (Excellent/Average/High), not
    # reachability. A numeric latency is the most reliable indicator that a host replied.
    online = [
        h for h in hosts
        if h.get("latency_ms") is not None
        or str(h.get("status", "")).strip().lower() in ("online", "up", "reachable")
    ]
    latencies = []
    for row in hosts:
        try:
            if row.get("latency_ms") is not None: latencies.append(float(row["latency_ms"]))
        except (TypeError, ValueError): pass
    opened = [p for p in ports if str(p.get("state", "")).lower() == "open"]
    good = [m for m in monitor if m.get("online")]
    loss = (100 * (len(monitor) - len(good)) / len(monitor)) if monitor else None
    jitter_samples = []
    monitor_latencies = []
    for row in monitor:
        try:
            if row.get("online") and row.get("latency_ms") is not None: monitor_latencies.append(float(row["latency_ms"]))
        except (TypeError, ValueError): pass
    jitter_samples = [abs(b-a) for a,b in zip(monitor_latencies, monitor_latencies[1:])]
    alerts = [dict(r) for r in getattr(app, "alert_events", [])]
    availability = (100.0 * len(good) / len(monitor)) if monitor else None
    return {
        "application": APP_TITLE, "version": "2.10.0", "generated_at": now.strftime("%Y-%m-%d %H:%M:%S %Z"),
        "local_ipv4": getattr(app, "local_ip_value", None) or (app.ip_var.get().strip() if hasattr(app, "ip_var") else ""),
        "public_ip": getattr(app, "public_ip_value", None) or "Unavailable",
        "summary": {"hosts_total": len(hosts), "hosts_online": len(online), "hosts_no_response": max(0, len(hosts)-len(online)),
                    "open_ports": len(opened), "port_records": len(ports), "dns_records": len(dns),
                    "interfaces": len(interfaces), "monitor_samples": len(monitor), "packet_loss_pct": loss,
                    "availability_pct": availability, "alert_events": len(alerts),
                    "avg_latency_ms": statistics.mean(latencies) if latencies else None,
                    "min_latency_ms": min(latencies) if latencies else None, "max_latency_ms": max(latencies) if latencies else None,
                    "avg_jitter_ms": statistics.mean(jitter_samples) if jitter_samples else None},
        "hosts": hosts, "ports": ports, "dns": dns, "interfaces": interfaces, "monitor": monitor,
        "history": list(getattr(app, "scan_snapshots", [])), "alerts": alerts,
    }


def _rows(payload):
    return [
        ("HOST INVENTORY", ["IP Address", "Hostname", "Latency (ms)", "Status", "Timestamp"],
         [[r.get("ip"), r.get("hostname"), r.get("latency_ms"), r.get("status"), r.get("timestamp")] for r in payload["hosts"]]),
        ("PORT ASSESSMENT", ["Target", "Resolved IP", "Port", "Service", "State"],
         [[r.get("host"), r.get("ip"), r.get("port"), r.get("service"), r.get("state")] for r in payload["ports"]]),
        ("DNS LOOKUPS", ["Record Type", "Name", "Value / Response"],
         [[r.get("type"), r.get("name"), r.get("value")] for r in payload["dns"]]),
        ("NETWORK INTERFACES", ["Name", "Address", "Netmask", "Status", "Details"],
         [[r.get("name"), r.get("address"), r.get("netmask"), r.get("status"), json.dumps(r, ensure_ascii=False)] for r in payload["interfaces"]]),
        ("PING MONITOR", ["Timestamp", "Target", "Availability", "Latency (ms)"],
         [[r.get("timestamp"), r.get("target"), "ONLINE" if r.get("online") else "NO RESPONSE", r.get("latency_ms")] for r in payload["monitor"]]),
        ("ALERT EVENTS", ["Timestamp", "Target", "Event", "Previous State", "Current State", "Latency (ms)"],
         [[r.get("timestamp"), r.get("target"), r.get("event"), r.get("previous_state"), r.get("current_state"), r.get("latency_ms")] for r in payload.get("alerts", [])]),
    ]


def export_json(payload, path):
    with open(path, "w", encoding="utf-8") as f: json.dump(payload, f, indent=2, ensure_ascii=False, default=str)


def export_csv(payload, path):
    fields = ["ip", "hostname", "latency_ms", "status", "timestamp"]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore"); writer.writeheader(); writer.writerows(payload["hosts"])


def export_markdown(payload, path):
    s = payload["summary"]
    lines = [f"# {APP_TITLE}", "", f"**Generated:** {payload['generated_at']}  ", f"**Version:** {payload['version']}  ", f"**Local IPv4:** {_safe(payload['local_ipv4'])}  ", f"**Public IP:** {_safe(payload['public_ip'])}", "", "## Executive Summary", "", "| Metric | Value |", "|---|---:|"]
    for label, key in [("Discovered Hosts","hosts_total"),("Responsive Hosts","hosts_online"),("No Response","hosts_no_response"),("Open TCP Ports","open_ports"),("DNS Records","dns_records"),("Monitor Samples","monitor_samples"),("Availability (%)","availability_pct"),("Alert Events","alert_events"),("Packet Loss","packet_loss_pct"),("Average Latency (ms)","avg_latency_ms")]:
        value=s.get(key); lines.append(f"| {label} | {_safe(round(value, 2) if isinstance(value, (int,float)) else value)}{'%' if key=='packet_loss_pct' and value is not None else ''} |")
    for title, headers, rows in _rows(payload):
        lines += ["", f"## {title.title()}", "", "| " + " | ".join(headers) + " |", "|" + "|".join(["---"]*len(headers)) + "|"]
        for row in rows:
            lines.append("| " + " | ".join(str(_safe(v)).replace("|", "\\|").replace("\n", " ") for v in row) + " |")
        if not rows: lines.append("| No data available |" + " |"*(len(headers)-1))
    lines += ["", "---", "*Generated locally by Network Operations Console. Findings reflect observed data only and are not a security certification.*", ""]
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def export_text(payload, path):
    s=payload["summary"]
    lines=[APP_TITLE, "="*72, f"Generated: {payload['generated_at']}", f"Version: {payload['version']}", f"Local IPv4: {_safe(payload['local_ipv4'])}", f"Public IP: {_safe(payload['public_ip'])}", "", "EXECUTIVE SUMMARY", "-"*72]
    for label,key in [("Hosts discovered","hosts_total"),("Responsive hosts","hosts_online"),("Hosts with no response","hosts_no_response"),("Open ports","open_ports"),("DNS records","dns_records"),("Monitor samples","monitor_samples"),("Availability (%)","availability_pct"),("Alert events","alert_events"),("Packet loss (%)","packet_loss_pct"),("Average latency (ms)","avg_latency_ms"),("Average jitter (ms)","avg_jitter_ms")]: lines.append(f"{label:28} {_safe(round(s[key],2) if isinstance(s.get(key),(int,float)) else s.get(key))}")
    for title, headers, rows in _rows(payload):
        lines += ["", title, "-"*72]
        lines.append(" | ".join(headers))
        for row in rows: lines.append(" | ".join(_safe(v) for v in row))
        if not rows: lines.append("No data available.")
    lines += ["", "CONFIDENTIALITY: Contains network inventory and diagnostic data. Handle according to organizational policy."]
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def export_html(payload, path):
    s=payload["summary"]
    def e(v): return html.escape(_safe(v))
    cards=[("Hosts Discovered",s["hosts_total"]),("Responsive Hosts",s["hosts_online"]),("Open TCP Ports",s["open_ports"]),("Packet Loss",(f"{s['packet_loss_pct']:.1f}%" if s["packet_loss_pct"] is not None else "N/A")),("Avg. Latency",(f"{s['avg_latency_ms']:.2f} ms" if s["avg_latency_ms"] is not None else "N/A")),("Availability",(f"{s['availability_pct']:.1f}%" if s.get("availability_pct") is not None else "N/A")),("Alert Events",s.get("alert_events", 0))]
    sections=[]
    for title,headers,rows in _rows(payload):
        head="".join(f"<th>{e(x)}</th>" for x in headers)
        body="".join("<tr>"+"".join(f"<td>{e(v)}</td>" for v in row)+"</tr>" for row in rows) or f'<tr><td colspan="{len(headers)}" class="empty">No data collected for this section.</td></tr>'
        sections.append(f'<section><div class="section-head"><h2>{e(title)}</h2><span>{len(rows):,} RECORDS</span></div><div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div></section>')
    card_html="".join(f'<div class="metric"><span>{e(k)}</span><strong>{e(v)}</strong></div>' for k,v in cards)
    doc=f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Network Operations | Assessment Report</title><style>
:root{{--navy:#14233b;--blue:#2457a7;--teal:#0f766e;--ink:#1c2b40;--muted:#607086;--line:#dce5f0;--paper:#fff;--bg:#eef3f9}}*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);font:14px/1.5 'Segoe UI',Arial,sans-serif;color:var(--ink)}}.shell{{max-width:1440px;margin:auto;padding:32px}}header{{background:linear-gradient(120deg,#14233b,#1d3c66);color:#fff;padding:34px;border-radius:18px;box-shadow:0 12px 30px #14233b22}}.eyebrow{{letter-spacing:2px;font-size:11px;color:#9cc5ff;font-weight:700}}h1{{font-size:30px;margin:8px 0}}.subtitle{{color:#c8d7ea}}.meta{{display:flex;gap:22px;flex-wrap:wrap;margin-top:22px;color:#dce7f6;font-size:12px}}.metrics{{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:12px;margin:18px 0}}.metric{{background:#fff;padding:18px;border:1px solid var(--line);border-radius:12px;min-height:98px}}.metric span{{display:block;color:var(--muted);font-size:11px;text-transform:uppercase;letter-spacing:.7px}}.metric strong{{display:block;font-size:25px;margin-top:7px;color:var(--navy);overflow-wrap:anywhere}}section{{background:#fff;border:1px solid var(--line);border-radius:14px;margin:14px 0;overflow:hidden}}.section-head{{display:flex;align-items:center;justify-content:space-between;padding:18px 20px;border-bottom:1px solid var(--line)}}h2{{font-size:15px;letter-spacing:.5px;margin:0}}.section-head span{{font-size:10px;letter-spacing:1px;color:var(--muted)}}.table-wrap{{overflow:auto}}table{{border-collapse:collapse;width:100%;font-size:12px}}th{{text-align:left;background:#f4f7fb;color:#4c607a;text-transform:uppercase;letter-spacing:.45px;font-size:10px}}th,td{{padding:11px 13px;border-bottom:1px solid #e8edf4;vertical-align:top;white-space:pre-wrap;overflow-wrap:anywhere}}tbody tr:nth-child(even){{background:#fbfcfe}}.empty{{color:var(--muted);text-align:center;padding:24px}}footer{{padding:20px 4px;color:var(--muted);font-size:11px}}@media(max-width:1000px){{.metrics{{grid-template-columns:repeat(3,1fr)}}.shell{{padding:16px}}}}@media print{{body{{background:white}}.shell{{padding:0;max-width:none}}header{{box-shadow:none;print-color-adjust:exact}}section,.metric{{break-inside:avoid;box-shadow:none}}.metrics{{grid-template-columns:repeat(3,1fr)}}}}
</style></head><body><main class="shell"><header><div class="eyebrow">IT OPERATIONS / NETWORK INTELLIGENCE</div><h1>Network Assessment Report</h1><div class="subtitle">Executive overview and technical evidence from the latest available session data.</div><div class="meta"><span>GENERATED: {e(payload['generated_at'])}</span><span>VERSION: {e(payload['version'])}</span><span>LOCAL IP: {e(payload['local_ipv4'])}</span><span>PUBLIC IP: {e(payload['public_ip'])}</span></div></header><div class="metrics">{card_html}</div>{''.join(sections)}<footer>INTERNAL USE · This report describes observed data only. ICMP non-response does not conclusively establish host downtime. Open ports require contextual review and do not alone prove a vulnerability.<br>Generated locally by Network Operations Console.</footer></main></body></html>'''
    Path(path).write_text(doc, encoding="utf-8")


def export_xlsx(payload, path):
    try: from openpyxl import Workbook
    except ImportError as exc: raise RuntimeError("Excel export requires openpyxl. Run: python -m pip install openpyxl") from exc
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb=Workbook(); ws=wb.active; ws.title="Executive Summary"; ws.sheet_view.showGridLines=False
    ws.merge_cells("A1:D1"); ws["A1"]=APP_TITLE; ws["A1"].font=Font(name="Aptos Display",size=20,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=NAVY); ws["A1"].alignment=Alignment(vertical="center"); ws.row_dimensions[1].height=38
    ws.merge_cells("A2:D2"); ws["A2"]="NETWORK ASSESSMENT | EXECUTIVE BRIEF"; ws["A2"].font=Font(size=10,bold=True,color="FFFFFF"); ws["A2"].fill=PatternFill("solid",fgColor=BLUE); ws.row_dimensions[2].height=24
    meta=[("Generated At",payload["generated_at"]),("Application Version",payload["version"]),("Local IPv4",_safe(payload["local_ipv4"])),("Public IP",_safe(payload["public_ip"]))]
    for i,(k,v) in enumerate(meta,4): ws.cell(i,1,k);ws.cell(i,2,v);ws.cell(i,1).font=Font(bold=True,color=NAVY);ws.cell(i,1).fill=PatternFill("solid",fgColor=LIGHT);ws.merge_cells(start_row=i,start_column=2,end_row=i,end_column=4)
    ws.cell(9,1,"KEY PERFORMANCE INDICATORS");ws.merge_cells("A9:D9");ws["A9"].font=Font(bold=True,color="FFFFFF");ws["A9"].fill=PatternFill("solid",fgColor=TEAL)
    ws.append(["Metric","Value","Metric","Value"])
    s=payload["summary"]; kpis=[("Hosts Discovered",s["hosts_total"]),("Responsive Hosts",s["hosts_online"]),("No Response",s["hosts_no_response"]),("Open TCP Ports",s["open_ports"]),("DNS Records",s["dns_records"]),("Network Interfaces",s["interfaces"]),("Monitor Samples",s["monitor_samples"]),("Packet Loss (%)",round(s["packet_loss_pct"],2) if s["packet_loss_pct"] is not None else "N/A"),("Average Latency (ms)",round(s["avg_latency_ms"],2) if s["avg_latency_ms"] is not None else "N/A"),("Minimum Latency (ms)",round(s["min_latency_ms"],2) if s["min_latency_ms"] is not None else "N/A"),("Maximum Latency (ms)",round(s["max_latency_ms"],2) if s["max_latency_ms"] is not None else "N/A"),("Average Jitter (ms)",round(s["avg_jitter_ms"],2) if s["avg_jitter_ms"] is not None else "N/A")]
    for i in range(0,len(kpis),2): ws.append([kpis[i][0],kpis[i][1],kpis[i+1][0] if i+1<len(kpis) else "",kpis[i+1][1] if i+1<len(kpis) else ""])
    for row in ws.iter_rows(min_row=10,max_row=ws.max_row,min_col=1,max_col=4):
        for c in row: c.alignment=Alignment(vertical="center",wrap_text=True); c.border=Border(bottom=Side(style="hair",color="DCE5F0"))
    for c in ws[10]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor=BLUE)
    for col,width in {"A":30,"B":28,"C":30,"D":28}.items(): ws.column_dimensions[col].width=width
    ws.freeze_panes="A10"; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.page_margins.left=.25;ws.page_margins.right=.25
    for title,headers,rows in _rows(payload):
        name={"HOST INVENTORY":"Host Inventory","PORT ASSESSMENT":"Port Assessment","DNS LOOKUPS":"DNS Lookups","NETWORK INTERFACES":"Interfaces","PING MONITOR":"Ping Monitor","ALERT EVENTS":"Alert Events"}[title]
        sh=wb.create_sheet(name);sh.sheet_view.showGridLines=False;sh.freeze_panes="A2";sh.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{max(1,len(rows)+1)}"
        sh.append(headers)
        for row in rows: sh.append([_safe(v) for v in row])
        for c in sh[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor=NAVY);c.alignment=Alignment(wrap_text=True,vertical="center")
        sh.row_dimensions[1].height=28
        for row in sh.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True);c.border=Border(bottom=Side(style="hair",color="E1E8F0"))
        for idx,header in enumerate(headers,1):
            maxlen=len(str(header))
            for cell in list(sh.columns)[idx-1][1: min(sh.max_row,250)]: maxlen=max(maxlen,min(42,len(str(cell.value or ""))))
            sh.column_dimensions[get_column_letter(idx)].width=min(44,max(14,maxlen+2))
        if rows:
            ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
            table=Table(displayName="T"+"".join(ch for ch in name if ch.isalnum()),ref=ref);table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False);sh.add_table(table)
        sh.sheet_properties.pageSetUpPr.fitToPage=True;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0;sh.page_margins.left=.25;sh.page_margins.right=.25
    wb.properties.title="Network Assessment Report";wb.properties.subject="Network inventory and diagnostics";wb.properties.creator=APP_TITLE;wb.properties.description=f"Generated {payload['generated_at']}"
    wb.save(path)


def export_pdf(payload, path, executive_only=False):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
    except ImportError as exc: raise RuntimeError("PDF export requires reportlab. Run: python -m pip install reportlab") from exc
    page_size=landscape(A4) if not executive_only else A4
    doc=SimpleDocTemplate(path,pagesize=page_size,rightMargin=15*mm,leftMargin=15*mm,topMargin=18*mm,bottomMargin=16*mm,title="Network Assessment Report",author=APP_TITLE,subject="Network operations")
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name="CoverTitle",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=21,leading=25,textColor=colors.HexColor("#14233B"),alignment=TA_LEFT,spaceAfter=7)); styles.add(ParagraphStyle(name="Section",parent=styles["Heading2"],fontSize=12,leading=15,textColor=colors.HexColor("#2457A7"),spaceBefore=12,spaceAfter=7)); styles.add(ParagraphStyle(name="SmallCell",parent=styles["BodyText"],fontSize=7,leading=9,wordWrap="CJK")); styles.add(ParagraphStyle(name="Meta",parent=styles["BodyText"],fontSize=8,textColor=colors.HexColor("#607086"),leading=12))
    story=[Paragraph(APP_TITLE,styles["CoverTitle"]),Paragraph("NETWORK ASSESSMENT REPORT" if not executive_only else "EXECUTIVE NETWORK BRIEF",styles["Heading2"]),Paragraph(f"Generated: {html.escape(payload['generated_at'])} &nbsp;&nbsp;|&nbsp;&nbsp; Version: {payload['version']}<br/>Local IPv4: {html.escape(_safe(payload['local_ipv4']))} &nbsp;&nbsp;|&nbsp;&nbsp; Public IP: {html.escape(_safe(payload['public_ip']))}",styles["Meta"]),Spacer(1,8*mm)]
    s=payload["summary"]
    kpis=[["DISCOVERED HOSTS",str(s["hosts_total"]),"RESPONSIVE HOSTS",str(s["hosts_online"])],["OPEN TCP PORTS",str(s["open_ports"]),"DNS RECORDS",str(s["dns_records"])],["PACKET LOSS",f"{s['packet_loss_pct']:.1f}%" if s["packet_loss_pct"] is not None else "N/A","AVG. LATENCY",f"{s['avg_latency_ms']:.2f} ms" if s["avg_latency_ms"] is not None else "N/A"],["AVAILABILITY",f"{s['availability_pct']:.1f}%" if s.get("availability_pct") is not None else "N/A","ALERT EVENTS",str(s.get("alert_events",0))]]
    kt=Table(kpis,colWidths=[doc.width*.24,doc.width*.26,doc.width*.24,doc.width*.26]);kt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#EEF3F9")),("TEXTCOLOR",(0,0),(-1,-1),colors.HexColor("#14233B")),("FONTNAME",(0,0),(-1,-1),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),9),("GRID",(0,0),(-1,-1),.5,colors.HexColor("#DCE5F0")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),10),("BOTTOMPADDING",(0,0),(-1,-1),10)]));story += [kt,Spacer(1,6*mm)]
    if executive_only:
        story += [Paragraph("KEY OBSERVATIONS",styles["Section"]),Paragraph(f"The current session contains {s['hosts_total']} discovered host record(s), {s['open_ports']} open-port record(s), and {s['dns_records']} DNS record(s). " + (f"The monitor sample set shows {s['packet_loss_pct']:.1f}% no-response samples." if s["packet_loss_pct"] is not None else "No continuous-monitor sample set is available for packet-loss analysis."),styles["BodyText"]),Spacer(1,4*mm),Paragraph("INTERPRETATION NOTES",styles["Section"]),Paragraph("ICMP non-response does not conclusively prove a device is offline. Open ports are observations, not proof of a vulnerability. Review scope, firewall policy, and service ownership before making operational decisions.",styles["BodyText"])]
    else:
        for title,headers,rows in _rows(payload):
            story.append(Paragraph(title,styles["Section"]))
            if not rows: story.append(Paragraph("No data collected for this section.",styles["Meta"]));continue
            data=[[Paragraph(html.escape(str(h)),styles["SmallCell"]) for h in headers]]
            for row in rows:
                data.append([Paragraph(html.escape(_safe(v)).replace("\n","<br/>"),styles["SmallCell"]) for v in row])
            col_width=doc.width/len(headers); table=Table(data,colWidths=[col_width]*len(headers),repeatRows=1,hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#14233B")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#DCE5F0")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F6F8FB")]),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]));story.append(table)
    def footer(canvas, document):
        canvas.saveState(); w,h=page_size; canvas.setStrokeColor(colors.HexColor("#DCE5F0"));canvas.line(15*mm,11*mm,w-15*mm,11*mm);canvas.setFont("Helvetica",7);canvas.setFillColor(colors.HexColor("#607086"));canvas.drawString(15*mm,7*mm,"INTERNAL USE  |  NETWORK OPERATIONS CONSOLE");canvas.drawRightString(w-15*mm,7*mm,f"Page {document.page}");canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def export_bundle(payload, path):
    import tempfile
    with tempfile.TemporaryDirectory(prefix="noc_reports_") as temp:
        base=Path(temp); files=[]
        for name,func,ext in [("Full Technical Report",lambda p:export_pdf(payload,p,False),"pdf"),("Executive Brief",lambda p:export_pdf(payload,p,True),"pdf"),("Excel Workbook",lambda p:export_xlsx(payload,p),"xlsx"),("Interactive HTML Report",lambda p:export_html(payload,p),"html"),("Structured Data",lambda p:export_json(payload,p),"json"),("Host Inventory CSV",lambda p:export_csv(payload,p),"csv"),("Markdown Report",lambda p:export_markdown(payload,p),"md"),("Plain Text Report",lambda p:export_text(payload,p),"txt")]:
            dest=base/(name.replace(" ","_")+"."+ext);func(str(dest));files.append((dest,dest.name))
        with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
            for source,arc in files:z.write(source,arcname=arc)
            readme=base/"README.txt";readme.write_text("NETWORK OPERATIONS CONSOLE | REPORT PACK\n\nAll reports are generated locally from data available in the application session. Missing sections indicate that no data was collected for that category.\n\nFormats: PDF executive brief, PDF technical report, styled Excel workbook, HTML dashboard, JSON, CSV, Markdown and plain text.\n",encoding="utf-8")
            z.write(readme,arcname="README.txt")
